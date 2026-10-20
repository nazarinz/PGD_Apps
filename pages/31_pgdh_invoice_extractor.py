"""
pgdh_invoice_extractor.py
Ekstrak invoice PGDH (The Look Macao -> adidas) dari PDF ke Excel multi-sheet.

Pemakaian:
    python pgdh_invoice_extractor.py --input ./invoices --output ./output/PGDH_Invoices.xlsx
    python pgdh_invoice_extractor.py --input PGDH0810110089.pdf

Dependensi:
    pip install pdfplumber pandas openpyxl
"""
from __future__ import annotations

import argparse
import io
import logging
import os
import re
import zipfile
from concurrent.futures import ProcessPoolExecutor
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Callable

import pandas as pd
import pdfplumber

# --------------------------------------------------------------------------- #
# Config
# --------------------------------------------------------------------------- #
LOG = logging.getLogger("pgdh")

DATE = r"\d{4}-\d{2}-\d{2}"
NUM = r"[\d,]+(?:\.\d+)?"
INCOTERMS = ("EXW", "FCA", "FAS", "FOB", "CFR", "CIF", "CPT", "CIP", "DAP", "DPU", "DDP", "CNF")
# Kode ukuran fraksi di invoice: 3713 -> 37 1/3, 3913 -> 39 1/3, dst.
SIZE_FRACTIONS = {"13": " 1/3", "23": " 2/3", "12": " 1/2"}
TOLERANCE = 0.011  # toleransi pembulatan untuk validasi nilai uang

TOTAL_LABELS = {
    "Total_Quantity": rf"Total Quantity\s+({NUM})\s*$",
    "Total_Carton": rf"Total Carton\s+({NUM})",
    "Gross_Weight_KG": rf"Total Gross Weight\s+({NUM})",
    "Net_Weight_KG": rf"Total Net Weight\s+({NUM})",
    "Net_Net_Weight_KG": rf"Total Net Net Weight\s+({NUM})",
    "Total_PO_Net_Amount": rf"Total PO Net Amount\s+({NUM})",
    "Additional_Charge": rf"Additional charge\s+({NUM})",
    "Total_VAT": rf"Total VAT\s+({NUM})",
    "Invoice_Total": rf"Invoice Total\s+({NUM})",
}


# --------------------------------------------------------------------------- #
# Helpers
# --------------------------------------------------------------------------- #
def to_num(value: str | None) -> float | None:
    if value is None:
        return None
    try:
        return float(value.replace(",", ""))
    except ValueError:
        return None


def search(pattern: str, text: str, group: int = 1, flags: int = re.M) -> str | None:
    m = re.search(pattern, text, flags)
    return m.group(group).strip() if m else None


def normalize_size(code: str) -> str:
    """'3713' -> '37 1/3'; '8-' -> '8 1/2' (tanda '-' = setengah size); '13K' tetap."""
    if len(code) == 4 and code.isdigit() and code[2:] in SIZE_FRACTIONS:
        return code[:2] + SIZE_FRACTIONS[code[2:]]
    if code.endswith("-") and code[:-1].isdigit():
        return f"{code[:-1]} 1/2"
    return code


@dataclass
class Invoice:
    source_file: str
    header: dict[str, Any] = field(default_factory=dict)
    items: list[dict[str, Any]] = field(default_factory=list)
    sizes: list[dict[str, Any]] = field(default_factory=list)
    adjustments: list[dict[str, Any]] = field(default_factory=list)
    hts: list[dict[str, Any]] = field(default_factory=list)
    totals: dict[str, Any] = field(default_factory=dict)
    checks: list[dict[str, Any]] = field(default_factory=list)

    @property
    def invoice_no(self) -> str | None:
        return self.header.get("Invoice_Number")


# --------------------------------------------------------------------------- #
# PDF reading
# --------------------------------------------------------------------------- #
def read_pdf(pdf: "pdfplumber.PDF") -> tuple[str, dict[str, str]]:
    """Return (full_text, parties) -- parties diambil dengan crop kolom
    karena blok For Account Of / Manufacturer / Consignee berdampingan."""
    if not pdf.pages:
        raise ValueError("PDF tanpa halaman")
    text = "\n".join(page.extract_text() or "" for page in pdf.pages)
    parties = _extract_parties(pdf.pages[0])
    if "Invoice Number" not in text:
        raise ValueError("Bukan invoice PGDH / PDF tanpa text layer (scan?)")
    return text, parties


def _extract_parties(page: "pdfplumber.page.Page") -> dict[str, str]:
    words = page.extract_words()

    def first(txt: str, after: float = 0.0):
        return next((w for w in words if w["text"] == txt and w["top"] > after), None)

    w_for, w_mfr, w_con = first("For"), first("Actual"), first("Consignee")
    if not (w_for and w_mfr and w_con):
        return {}
    y0 = w_for["top"]
    stops = [w for w in words if w["text"] in ("Company", "GPS") and w["top"] > y0]
    stop = min(stops, key=lambda w: w["top"]) if stops else None
    y1 = stop["top"] - 2 if stop else page.height
    cols = {
        "For_Account_Of": (w_for["x0"], w_mfr["x0"] - 2),
        "Actual_Manufacturer": (w_mfr["x0"] - 2, w_con["x0"] - 2),
        "Consignee": (w_con["x0"] - 2, page.width),
    }
    out: dict[str, str] = {}
    for key, (xa, xb) in cols.items():
        txt = page.crop((xa, y0, xb, y1)).extract_text() or ""
        lines = [ln for ln in txt.splitlines()[1:] if ln.strip()]  # buang baris label
        out[key] = "\n".join(lines)
    return out


def extract_adjustment_columns(pdf: "pdfplumber.PDF") -> list[dict[str, str]]:
    """Ambil PO Lines / Customer Size / Sourcing Size tiap baris adjustment.

    Geometri kolom berbeda di setiap tabel, jadi tidak memakai posisi header. Pada baris
    yang memuat kode Zxxx dan 'USD', tiga kata di kirinya = ketiga kolom itu. Potongan
    sel di baris bawahnya (x0 sama dengan kolomnya) disambung tanpa spasi.
    Hasil berurutan sesuai urutan dokumen."""
    rows: list[dict[str, str]] = []
    if True:
        for page in pdf.pages:
            words = page.extract_words()
            anchors = []
            for w in words:
                if not re.fullmatch(r"Z[A-Z0-9]{3}", w["text"]):
                    continue
                line = [x for x in words if abs(x["top"] - w["top"]) < 2]
                left = sorted((x for x in line if x["x1"] <= w["x0"] + 0.5), key=lambda x: x["x0"])
                if len(left) >= 3 and any(x["text"] == "USD" for x in line):
                    anchors.append((w, left[-3:]))
            anchors.sort(key=lambda t: t[0]["top"])
            stops = [w["top"] for i, w in enumerate(words)
                     if w["text"] in ("Net", "Goods", "Page", "Total")
                     or (w["text"] == "PO" and i + 1 < len(words) and words[i + 1]["text"] == "No")]
            for k, (a_w, left) in enumerate(anchors):
                nxt = [t for t in stops if t > a_w["top"] + 1]
                if k + 1 < len(anchors):
                    nxt.append(anchors[k + 1][0]["top"] - 0.5)
                end = min(nxt) if nxt else 1e9
                cells = [left[i]["text"] for i in range(3)]
                for x in sorted(words, key=lambda x: (x["top"], x["x0"])):
                    if a_w["top"] + 2 < x["top"] < end:
                        for i, col in enumerate(left):
                            if abs(x["x0"] - col["x0"]) <= 1.5:
                                cells[i] += x["text"]
                rows.append({"Code": a_w["text"], "PO_Lines": cells[0],
                             "Customer_Sizes": cells[1], "Sourcing_Sizes": cells[2]})
    return rows


# --------------------------------------------------------------------------- #
# Parsers
# --------------------------------------------------------------------------- #
def parse_header(text: str, parties: dict[str, str]) -> dict[str, Any]:
    h: dict[str, Any] = {}
    m = re.search(r"(PGDH\d+)\s+(\d{6,})", text)
    h["Invoice_Number"], h["Log_No"] = (m.group(1), m.group(2)) if m else (None, None)

    m = re.search(rf"Document Date Invoice Date[^\n]*\n[^\n]*?({DATE})\s+({DATE})", text)
    h["Document_Date"], h["Invoice_Date"] = (m.group(1), m.group(2)) if m else (None, None)

    first_lines = text.splitlines()[:2]
    if len(first_lines) == 2:
        l1 = re.split(r"\s+Invoice as of", first_lines[0])[0]
        l2 = re.split(r"\s+Invoice Number", first_lines[1])[0]
        h["Issuer"] = f"{l1} {l2}".strip()

    parties = dict(parties)
    fa_lines = parties.get("For_Account_Of", "").splitlines()
    if "As Buying Agent For" in fa_lines:  # blok agent berada di kolom yang sama
        i = fa_lines.index("As Buying Agent For")
        parties["For_Account_Of"] = "\n".join(fa_lines[:i])
        parties["Buying_Agent"] = "\n".join(fa_lines[i + 1:])
    for key, val in parties.items():
        lines = val.splitlines()
        h[key] = lines[0] if lines else None
        h[f"{key}_Detail"] = " | ".join(lines)
    h["Consignee_Tax_ID"] = search(r"(?:Customer VAT nr\.|RUC:)\s*(\w+)", parties.get("Consignee", ""))

    m = re.search(r"^(Company Code|GPS Customer Number)[^\n]*Origin From\n([^\n]+)", text, re.M)
    if m:
        toks = m.group(2).split()
        h["Company_Code" if m.group(1) == "Company Code" else "GPS_Customer_Number"] = toks[0]
        h["Customer_VAT_No"] = " ".join(toks[1:-1]) or None
        h["Origin_From"] = toks[-1]

    h["Shipment_Country_Destination_To"] = search(
        r"Shipment Country Destination To[^\n]*\n([^\n]+)", text
    )

    inc = "|".join(INCOTERMS)
    m = re.search(rf"Payment Method[^\n]*\n(.+?)\s+({inc})\b\s*-?\s*(.*?)\s+([A-Z]{{3}})$", text, re.M)
    if m:
        h["Payment_Method"], h["Incoterm"], h["Incoterm_Place"], h["Currency"] = (
            g.strip() for g in m.groups()
        )

    m = re.search(rf"FCR Date[^\n]*\n({DATE})\s+({DATE})\s+(.+)", text)
    if m:
        h["FCR_Date"], h["Ex_Factory_Date"], h["Ship_Mode"] = (g.strip() for g in m.groups())

    m = re.search(r"Forwarder VAT Rate (Plant Code|Brand)\n(.+?)\s+([\d.]+)\s+(\w+)$", text, re.M)
    if m:
        h["Forwarder"], h["VAT_Rate"] = m.group(2), m.group(3)
        h["Plant_Code" if m.group(1) == "Plant Code" else "Brand"] = m.group(4)

    m = re.search(r"Customer Number Brand\n(\d+)\s+(\w+)", text)
    if m:
        h["Customer_Number"], h["Brand"] = m.groups()

    m = re.search(r"The exporter\s+(.+?)\s+(ID[A-Z0-9]+)\s+of the products", text, re.S)
    if m:
        h["Exporter_Name"], h["Exporter_Reg_No"] = " ".join(m.group(1).split()), m.group(2)
    h["Origin_Criterion"] = search(r"origin criterion met is (\w+)", text)
    h["Invoice_Text"] = search(r"Invoice Text[^\n]*\n([^\n]+)", text)
    return h


ADJ_RE = re.compile(
    rf"(?P<code>Z[A-Z0-9]{{3}})\s+(?P<desc>.+?)\s+(?P<rate>-?\d+(?:\.\d+)?)\s+(?P<cur>[A-Z]{{3}})\s+"
    rf"(?:per\s+)?(?:Unit\s+)?(?P<qty>-?\d[\d,]*)\s+(?P<amt>-?{NUM})"
)
HTS_RE = re.compile(r"(?P<hts>\d{10,12})\s+(?:(?P<mat>.+?)\s+)?(?P<gender>[A-Z])\s+(?P<qty>[\d,]+)$")
SIZE_HDR_RE = re.compile(r"^Customer Size\s+(.+?)\s+Total Quantity\s+Price")


def parse_item_line(ln: str) -> dict[str, Any] | None:
    """Baris item. Format A (lama): PO Line WorkingNo Article Desc G Cat Retailer Qty Pref.
    Format B (baru): PO Line MarketPO SalesOrder WorkingNo Article Desc G Cat Qty."""
    t = ln.split()
    if len(t) < 9 or not (re.fullmatch(r"\d{10}", t[0]) and t[1].isdigit()):
        return None
    i = 2
    refs: list[str] = []
    while i < len(t) and t[i].isdigit():
        refs.append(t[i])
        i += 1
    if not refs or i >= len(t) - 4:
        return None
    article, rest = t[i], t[i + 1:]
    pref = retailer = None
    if rest[-1] in ("Y", "N"):
        pref, qty, retailer, cat, gender = rest[-1], rest[-2], rest[-3], rest[-4], rest[-5]
        desc = rest[:-5]
    else:
        qty, cat, gender = rest[-1], rest[-2], rest[-3]
        desc = rest[:-3]
    if not (re.fullmatch(r"[\d,]+", qty) and re.fullmatch(r"[A-Z]", gender) and desc):
        return None
    market_po = sales_order = None
    if len(refs) >= 3:
        market_po, sales_order = refs[-3], refs[-2]
    elif len(refs) == 2:
        sales_order = refs[0]
    return {
        "PO_No": t[0],
        "PO_Line": int(t[1]),
        "Market_PO_Number": market_po,
        "Sales_Order_No": sales_order,
        "Working_No": refs[-1],
        "Article_No": article,
        "Article_Description": " ".join(desc),
        "Gender": gender,
        "Category": cat,
        "Retailer_Order_No": retailer,
        "Total_Qty": to_num(qty),
        "Preference_Indicator": pref,
    }


def parse_body(text: str, inv_no: str) -> tuple[list, list, list, list]:
    """State machine per item: item -> size/qty -> adjustments -> net amount -> HTS."""
    items: list[dict] = []
    sizes: list[dict] = []
    adjs: list[dict] = []
    hts: list[dict] = []
    cur: dict | None = None
    size_codes: list[str] = []
    section: str | None = None
    adj_buf: list[str] = []

    def key() -> dict[str, Any]:
        return {"Invoice_Number": inv_no, "PO_No": cur["PO_No"], "PO_Line": cur["PO_Line"]}

    for raw in text.splitlines():
        ln = raw.strip()
        item = parse_item_line(ln)
        if item:
            cur = {"Invoice_Number": inv_no, **item}
            items.append(cur)
            section, size_codes, adj_buf = None, [], []
            continue
        if cur is None or not ln:
            continue
        if ln.startswith("Total Quantity ") and section == "hts":
            section = None
            continue
        m = SIZE_HDR_RE.match(ln)
        if m:
            size_codes = m.group(1).split()
            continue
        m = re.match(rf"QTY\s+((?:{NUM}\s+)+{NUM})$", ln)
        if m and size_codes:
            nums = m.group(1).split()
            n = len(size_codes)
            if len(nums) != n + 3:
                LOG.warning("%s PO %s: jumlah size (%d) != jumlah qty (%d)", inv_no, cur["PO_No"], n, len(nums) - 3)
            for code, q in zip(size_codes, nums[:n]):
                sizes.append(
                    {**key(), "Article_No": cur["Article_No"], "Customer_Size_Code": code,
                     "Customer_Size": normalize_size(code), "Qty": to_num(q)}
                )
            tail = nums[n:]
            cur["Unit_Price"], cur["Item_Total_Amount"] = to_num(tail[1]), to_num(tail[2])
            size_codes = []
            continue
        if ln.startswith("Adjustments"):
            section, adj_buf = "adj", []
        elif section == "adj" and ln == "Net Amount":
            for m in ADJ_RE.finditer(" ".join(adj_buf)):
                adjs.append(
                    {**key(), "PO_Lines": None, "Customer_Sizes": None, "Sourcing_Sizes": None,
                     "Code": m["code"], "Description": m["desc"],
                     "Amount_Per_Unit": f'{m["rate"]} {m["cur"]} per Unit',
                     "Rate": to_num(m["rate"]), "Currency": m["cur"],
                     "Total_Quantity": to_num(m["qty"]), "Amount_Per_Charge": to_num(m["amt"]),
                     "Item_Net_Amount": None}
                )
            section = "net_wait"
        elif section == "adj":
            adj_buf.append(ln)
        elif section == "net_wait" and re.fullmatch(NUM, ln):
            cur["Item_Net_Amount"] = to_num(ln)
            section = None
        elif ln.startswith("Goods Description"):
            section = "hts"
        elif section == "hts":
            m = HTS_RE.search(ln)
            if m:
                cols = ln[: m.start()].split()
                hts.append(
                    {**key(), "Customer_Sizes": cols[0] if cols else None,
                     "Sourcing_Sizes": cols[1] if len(cols) > 1 else None, "HTS": m["hts"],
                     "Upper_Material": m["mat"], "Gender": m["gender"], "Quantity": to_num(m["qty"])}
                )
    nets = {(i["PO_No"], i["PO_Line"]): i.get("Item_Net_Amount") for i in items}
    for a in adjs:
        a["Item_Net_Amount"] = nets.get((a["PO_No"], a["PO_Line"]))
    return items, sizes, adjs, hts


def parse_totals(text: str, inv_no: str) -> dict[str, Any]:
    t: dict[str, Any] = {"Invoice_Number": inv_no}
    for key, pat in TOTAL_LABELS.items():
        t[key] = to_num(search(pat, text))
    return t


# --------------------------------------------------------------------------- #
# Validation
# --------------------------------------------------------------------------- #
def validate(inv: Invoice) -> list[dict[str, Any]]:
    def check(name: str, expected: float | None, actual: float | None, scope: str = "INVOICE") -> dict[str, Any]:
        ok = expected is not None and actual is not None and abs(expected - actual) <= TOLERANCE
        return {"Invoice_Number": inv.invoice_no, "Scope": scope, "Check": name,
                "Expected": expected, "Actual": actual, "Status": "OK" if ok else "MISMATCH"}

    t = inv.totals
    checks: list[dict[str, Any]] = []
    for it in inv.items:
        scope = f"PO {it['PO_No']}-{it['PO_Line']}"
        same = lambda rows: [r for r in rows if r["PO_No"] == it["PO_No"] and r["PO_Line"] == it["PO_Line"]]  # noqa: E731
        qty = it["Total_Qty"]
        checks += [
            check("Sum size qty = item qty", qty, sum(r["Qty"] or 0 for r in same(inv.sizes)), scope),
            check("Sum HTS qty = item qty", qty, sum(r["Quantity"] or 0 for r in same(inv.hts)), scope),
            check("Qty x price = item amount", it.get("Item_Total_Amount"),
                  round((qty or 0) * (it.get("Unit_Price") or 0), 2), scope),
            check("Item amount + adjustments = item net", it.get("Item_Net_Amount"),
                  round((it.get("Item_Total_Amount") or 0) + sum(r["Amount_Per_Charge"] or 0 for r in same(inv.adjustments)), 2), scope),
        ]
    checks += [
        check("Sum item qty = Total Quantity", t.get("Total_Quantity"), sum(i["Total_Qty"] or 0 for i in inv.items)),
        check("Sum item net = Total PO Net Amount", t.get("Total_PO_Net_Amount"),
              round(sum(i.get("Item_Net_Amount") or 0 for i in inv.items), 2)),
        check("PO Net + Additional charge + VAT = Invoice Total", t.get("Invoice_Total"),
              round((t.get("Total_PO_Net_Amount") or 0) + (t.get("Additional_Charge") or 0) + (t.get("Total_VAT") or 0), 2)),
    ]
    return checks


# --------------------------------------------------------------------------- #
# Orchestration
# --------------------------------------------------------------------------- #
def parse_invoice(path: Path) -> Invoice:
    with pdfplumber.open(path) as pdf:
        text, parties = read_pdf(pdf)
        header = parse_header(text, parties)
        inv_no = header.get("Invoice_Number")
        if not inv_no:
            raise ValueError("Invoice Number tidak ditemukan")
        inv = Invoice(source_file=path.name, header=header)
        inv.items, inv.sizes, inv.adjustments, inv.hts = parse_body(text, inv_no)
        cols = extract_adjustment_columns(pdf)
    if len(cols) == len(inv.adjustments) and all(c["Code"] == a["Code"] for c, a in zip(cols, inv.adjustments)):
        for c, a in zip(cols, inv.adjustments):
            a.update({k: c[k] for k in ("PO_Lines", "Customer_Sizes", "Sourcing_Sizes")})
    else:
        LOG.warning("%s: kolom PO Lines/Size adjustment tidak bisa dipetakan (%d vs %d baris)",
                    path.name, len(cols), len(inv.adjustments))
    if not inv.items:
        LOG.warning("%s: tidak ada baris item terbaca", path.name)
    inv.totals = parse_totals(text, inv_no)
    inv.checks = validate(inv)
    return inv


def _safe_parse(path: Path) -> tuple[Path, Invoice | None, str | None]:
    """Dipakai worker: exception diubah jadi pesan supaya satu PDF rusak tidak menghentikan batch."""
    try:
        return path, parse_invoice(path), None
    except Exception as exc:  # noqa: BLE001
        return path, None, f"{type(exc).__name__}: {exc}"


def process_files(
    pdfs: list[Path],
    workers: int | None = None,
    progress: bool = True,
    on_progress: Callable[[int, int, str], None] | None = None,
) -> tuple[list[Invoice], list[dict]]:
    """Proses banyak PDF (paralel bila workers > 1). Invoice duplikat dilewati dan dilaporkan.
    Hasil selalu berurutan sesuai urutan input."""
    workers = workers or min(4, os.cpu_count() or 1)

    def collect(results) -> list:
        out = []
        for k, res in enumerate(_with_progress(results, len(pdfs), progress), 1):
            out.append(res)
            if on_progress:
                on_progress(k, len(pdfs), res[0].name)
        return out

    if workers > 1 and len(pdfs) > 1:
        with ProcessPoolExecutor(max_workers=workers) as ex:
            outcomes = collect(ex.map(_safe_parse, pdfs, chunksize=1))
    else:
        outcomes = collect(map(_safe_parse, pdfs))

    invoices: list[Invoice] = []
    issues: list[dict] = []
    seen: dict[str, str] = {}
    for path, inv, err in outcomes:
        if err:
            issues.append({"Source_File": path.name, "Type": "ERROR", "Detail": err})
        elif inv.invoice_no in seen:
            issues.append({"Source_File": path.name, "Type": "DUPLICATE",
                           "Detail": f"Invoice {inv.invoice_no} sudah ada dari file {seen[inv.invoice_no]}; dilewati"})
        else:
            seen[inv.invoice_no] = path.name
            invoices.append(inv)
    LOG.info("Selesai: %d invoice, %d masalah (%d duplikat)", len(invoices), len(issues),
             sum(i["Type"] == "DUPLICATE" for i in issues))
    return invoices, issues


def _with_progress(it, total: int, enabled: bool):
    if not enabled:
        return it
    try:
        from tqdm.auto import tqdm
        return tqdm(it, total=total, desc="Parsing PDF")
    except ImportError:
        return it


def collect_pdfs(src: Path) -> list[Path]:
    if src.is_file():
        return [src]
    return sorted(src.rglob("*.pdf"))


def build_frames(invoices: list[Invoice], issues: list[dict]) -> dict[str, pd.DataFrame]:
    def with_src(rows_attr: str) -> pd.DataFrame:
        rows = []
        for inv in invoices:
            for r in getattr(inv, rows_attr):
                rows.append({"Source_File": inv.source_file, **r})
        return pd.DataFrame(rows)

    frames = {
        "Header": pd.DataFrame([{"Source_File": i.source_file, **i.header} for i in invoices]),
        "Items": with_src("items"),
        "Size_Breakdown": with_src("sizes"),
        "Adjustments": with_src("adjustments"),
        "HTS": with_src("hts"),
        "Totals": pd.DataFrame([{"Source_File": i.source_file, **i.totals} for i in invoices]),
        "Validation": with_src("checks"),
    }
    if issues:
        frames["Issues"] = pd.DataFrame(issues)
    return frames


def write_excel(frames: dict[str, pd.DataFrame], out: Path) -> None:
    out.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(out, engine="openpyxl") as xw:
        for name, df in frames.items():
            df.to_excel(xw, sheet_name=name, index=False)
            ws = xw.sheets[name]
            ws.freeze_panes = "A2"
            for col in ws.columns:
                width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
                ws.column_dimensions[col[0].column_letter].width = min(width + 2, 60)
    LOG.info("Excel tersimpan: %s", out)


# --------------------------------------------------------------------------- #
# Upload handling (untuk Streamlit / web)
# --------------------------------------------------------------------------- #
MAX_ZIP_BYTES = 500 * 1024 * 1024


def save_uploads(files: list[tuple[str, bytes]], dest: Path) -> list[Path]:
    """Tulis hasil upload ke `dest`. File ZIP diekstrak (hanya .pdf; nama dibersihkan
    sehingga aman dari path traversal). Nama file asli dipertahankan untuk kolom Source_File."""
    paths: list[Path] = []

    def put(name: str, data: bytes) -> None:
        folder = dest / str(len(paths))
        folder.mkdir(parents=True, exist_ok=True)
        target = folder / Path(name).name
        target.write_bytes(data)
        paths.append(target)

    for name, data in files:
        if name.lower().endswith(".zip"):
            with zipfile.ZipFile(io.BytesIO(data)) as zf:
                if sum(i.file_size for i in zf.infolist()) > MAX_ZIP_BYTES:
                    raise ValueError(f"{name}: isi ZIP terlalu besar (> {MAX_ZIP_BYTES // 2**20} MB)")
                for info in zf.infolist():
                    if not info.is_dir() and info.filename.lower().endswith(".pdf"):
                        put(info.filename, zf.read(info))
        else:
            put(name, data)
    return paths


# --------------------------------------------------------------------------- #
# Summary adjustment per invoice (dipakai bila upload > 1 invoice)
# --------------------------------------------------------------------------- #
# (label di laporan, daftar Description yang dijumlahkan). Description lain yang tidak
# terdaftar tetap muncul sebagai baris sendiri, jadi Total selalu = jumlah semua adjustment.
ADJUSTMENT_SUMMARY_GROUPS: list[tuple[str, tuple[str, ...]]] = [
    ("SHAS", ("SHAS price",)),
    ("MOS quantity flat, PROMO", ("MOS quantity flat", "PROMO")),
]
SUMMARY_SHEET = "Summary"


def build_adjustment_summary(invoices: list[Invoice]) -> list[tuple[str, float | None, str]]:
    """Baris summary: (label, jumlah, jenis) dengan jenis 'header' | 'line' | 'total' | 'blank'.
    Jumlah = penjumlahan Amount_Per_Charge per kelompok Description."""
    rows: list[tuple[str, float | None, str]] = []
    for inv in invoices:
        sums: dict[str, float] = {}
        for a in inv.adjustments:
            sums[(a["Description"] or "").strip()] = sums.get((a["Description"] or "").strip(), 0.0) + (a["Amount_Per_Charge"] or 0.0)
        used: set[str] = set()
        rows.append((inv.invoice_no or "", None, "header"))
        total = 0.0
        for label, descs in ADJUSTMENT_SUMMARY_GROUPS:
            keys = {d.lower() for d in descs}
            amt = sum(v for k, v in sums.items() if k.lower() in keys)
            used |= {k for k in sums if k.lower() in keys}
            rows.append((label, round(amt, 2), "line"))
            total += amt
        for k, v in sums.items():
            if k not in used:
                rows.append((k, round(v, 2), "line"))
                total += v
        rows.append(("Total", round(total, 2), "total"))
        rows.append(("", None, "blank"))
    return rows[:-1] if rows else rows


def adjustment_summary_frame(invoices: list[Invoice]) -> pd.DataFrame:
    return pd.DataFrame([(l, a) for l, a, _ in build_adjustment_summary(invoices)], columns=["Keterangan", "Jumlah"])


def _write_summary_sheet(wb, invoices: list[Invoice]) -> None:
    from openpyxl.styles import Font

    ws = wb.create_sheet(SUMMARY_SHEET)
    bold = Font(bold=True)
    for r, (label, amount, kind) in enumerate(build_adjustment_summary(invoices), start=1):
        ws.cell(r, 1, label or None)
        if amount is not None:
            c = ws.cell(r, 2, amount)
            c.number_format = "#,##0.00"
        if kind in ("header", "total"):
            ws.cell(r, 1).font = bold
            ws.cell(r, 2).font = bold
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 14


# --------------------------------------------------------------------------- #
# Format laporan ringkas: Sheet 1 (item) + Sheet 2 (adjustment)
# --------------------------------------------------------------------------- #
SHEET1_COLS = ["PO_No", "PO_Line", "Market_PO_Number", "Sales_Order_No", "Working_No", "Article_No",
               "Total_Qty", "Preference_Indicator", "Unit_Price", "Item_Total_Amount", "Item_Net_Amount"]
SHEET2_COLS = ["Invoice_Number", "PO_No", "Description", "Amount_Per_Unit", "Rate", "Currency",
               "Total_Quantity", "Amount_Per_Charge", "Item_Net_Amount"]
EXTRA_SHEETS = {"Header": "Header", "Size_Breakdown": "Size_Breakdown", "HTS": "HTS",
                "Totals": "Totals", "Validation": "Validation", "Issues": "Issues",
                "Adjustments": "Adjustments_Detail"}
LAYOUT_SEPARATE, LAYOUT_SIDE_BY_SIDE = "separate", "side_by_side"


def _int_cols(df: pd.DataFrame, cols: list[str]) -> pd.DataFrame:
    for c in cols:
        if c in df:
            df[c] = pd.to_numeric(df[c], errors="coerce").astype("Int64")
    return df


def build_report_frames(invoices: list[Invoice], with_invoice_no: bool = False) -> tuple[pd.DataFrame, pd.DataFrame]:
    """Sheet 1 = satu baris per PO line; Sheet 2 = satu baris per adjustment."""
    items = pd.DataFrame([i for inv in invoices for i in inv.items])
    adjs = pd.DataFrame([a for inv in invoices for a in inv.adjustments])
    cols1 = (["Invoice_Number"] if with_invoice_no else []) + SHEET1_COLS
    sheet1 = items.reindex(columns=cols1)
    sheet2 = adjs.reindex(columns=SHEET2_COLS)
    return _int_cols(sheet1, ["PO_Line", "Total_Qty"]), _int_cols(sheet2, ["Total_Quantity"])


def _style_sheet(ws, min_col: int = 1) -> None:
    ws.freeze_panes = "A2"
    for col in ws.iter_cols(min_col=min_col):
        width = max((len(str(c.value)) for c in col if c.value is not None), default=0)
        ws.column_dimensions[col[0].column_letter].width = min(max(width + 2, 8), 60)


def report_excel_bytes(
    invoices: list[Invoice],
    issues: list[dict],
    *,
    with_invoice_no: bool = False,
    layout: str = LAYOUT_SEPARATE,
    extra_sheets: bool = False,
    summary: bool = True,
) -> bytes:
    """Bangun file Excel di memori. layout: 'separate' (Sheet 1 & Sheet 2) atau
    'side_by_side' (satu sheet, dua tabel berdampingan dengan 1 kolom kosong).
    Bila invoice > 1 dan summary=True, ditambah sheet 'Summary' (jumlah adjustment per invoice)."""
    sheet1, sheet2 = build_report_frames(invoices, with_invoice_no)
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        if layout == LAYOUT_SIDE_BY_SIDE:
            sheet1.to_excel(xw, sheet_name="Report", index=False, startcol=0)
            sheet2.to_excel(xw, sheet_name="Report", index=False, startcol=len(sheet1.columns) + 1)
            _style_sheet(xw.sheets["Report"])
        else:
            sheet1.to_excel(xw, sheet_name="Sheet 1", index=False)
            sheet2.to_excel(xw, sheet_name="Sheet 2", index=False)
            _style_sheet(xw.sheets["Sheet 1"])
            _style_sheet(xw.sheets["Sheet 2"])
        if summary and len(invoices) > 1:
            _write_summary_sheet(xw.book, invoices)
        if extra_sheets:
            frames = build_frames(invoices, issues)
            for key, name in EXTRA_SHEETS.items():
                if key in frames:
                    frames[key].to_excel(xw, sheet_name=name, index=False)
                    _style_sheet(xw.sheets[name])
    return buf.getvalue()


def main() -> int:
    ap = argparse.ArgumentParser(description="Ekstrak invoice PGDH PDF -> Excel")
    ap.add_argument("--input", required=True, type=Path, help="File PDF atau folder berisi PDF")
    ap.add_argument("--output", type=Path, default=Path("output/PGDH_Invoices.xlsx"))
    ap.add_argument("--workers", type=int, default=None, help="Jumlah proses paralel (default: min(4, CPU))")
    args = ap.parse_args()
    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")

    pdfs = collect_pdfs(args.input)
    if not pdfs:
        LOG.error("Tidak ada PDF ditemukan di %s", args.input)
        return 1

    invoices, issues = process_files(pdfs, workers=args.workers)
    if not invoices:
        LOG.error("Tidak ada invoice yang berhasil diparse")
        return 1
    bad = sum(c["Status"] != "OK" for i in invoices for c in i.checks)
    LOG.info("%d pengecekan validasi gagal", bad)
    write_excel(build_frames(invoices, issues), args.output)
    return 0 if not any(i["Type"] == "ERROR" for i in issues) else 2


if __name__ == "__main__":
    raise SystemExit(main())
