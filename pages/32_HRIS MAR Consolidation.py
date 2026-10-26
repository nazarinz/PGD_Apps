"""
HRIS Monthly Attendance ETL (Streamlit)
---------------------------------------
Upload PDF "Monthly Attendance Report" HRIS (PT.GLOSTAR INDONESIA) -> parse -> QA -> Excel.

Run:
    pip install streamlit pdfplumber pandas openpyxl
    streamlit run hris_attendance_etl_app.py

Output workbook: README, Attendance_Daily, Employee_Monthly, Dim_Employee, Dim_Status, QA_Checks
"""
import re
import io
import calendar
import datetime as dt
from collections import Counter

import pandas as pd
import pdfplumber
import streamlit as st
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

# ----------------------------------------------------------------------------
# Config
# ----------------------------------------------------------------------------
# Kolom ringkasan per karyawan di PDF (urut dari kiri ke kanan). Januari punya 18 kolom, Februari 19 (+SOR).
SUM_COLS = ['H', 'A', 'I', 'CT', 'S', 'SD', 'TL', 'IR', 'DL', 'CH', 'CHK', 'HD', 'PC', 'LC',
            'N', 'CL4', 'CL5', 'CL6', 'SOR']

STATUS = {  # kode: (deskripsi, grup)
    'N': ('Normal', 'Hadir'), 'TL': ('Terlambat', 'Hadir'), 'PC': ('Pulang Cepat', 'Hadir'),
    'DL': ('Dinas Luar', 'Hadir'),
    'CT': ('Cuti', 'Cuti'), 'CH': ('Cuti Hamil', 'Cuti'),
    'SD': ('Surat Dokter', 'Sakit'), 'S': ('Sakit', 'Sakit'),
    'I': ('Izin', 'Izin'), 'IR': ('Izin Resmi', 'Izin'),
    'A': ('Alpa', 'Alpa'), 'H': ('Hari Libur', 'Libur'),
    'OFF': ('Weekend (tidak ada kode di PDF)', 'Weekend'),
    'NODATA': ('Hari kerja tanpa kode di PDF', 'No Data'),
}
RAW_ALIAS = {'I*': 'IR'}  # di ringkasan HRIS, I* dihitung sebagai IR

# Batas koordinat (pt) untuk layout PDF ini
X_DAY_MIN, X_DAY_MAX = 150, 570       # area kolom tanggal
X_SUMMARY_MIN = 570                    # area kolom ringkasan
X_TOTAL_SPLIT = 568.5                  # pemisah kolom hari-31 dan kolom total lembur
X_NAME_MIN, X_NAME_MAX = 66, 300       # area nama
X_NAME_JUNK_SPACE = 82                 # spasi < x ini adalah sisa kolom NIK, bukan spasi nama


# ----------------------------------------------------------------------------
# Extract / Transform
# ----------------------------------------------------------------------------
def _words(page, cond):
    f = page.filter(lambda o: o.get('object_type') != 'char' or cond(o))
    return f.extract_words(x_tolerance=1.5, keep_blank_chars=False)


def parse_pdf(file_obj):
    """Parse satu PDF HRIS. Return (period 'YYYYMM', list[dict] per karyawan)."""
    out, period = [], None
    with pdfplumber.open(file_obj) as pdf:
        for pg in pdf.pages:
            m = re.search(r'Bulan\s*:\s*(\d{6})', pg.extract_text() or '')
            if m:
                period = m.group(1)
            chars = pg.chars

            # posisi header tanggal 1..31
            hw = _words(pg, lambda o: 70 < o['top'] < 76 and o['size'] < 8.6)
            day_x = {int(w['text']): (w['x0'] + w['x1']) / 2 for w in hw
                     if w['text'].isdigit() and 1 <= int(w['text']) <= 31 and X_DAY_MIN < w['x0'] < X_DAY_MAX}
            if len(day_x) < 31:       # halaman overflow (kolom lanjutan), lewati
                continue

            # anchor tiap karyawan = nomor urut (font 9) di kolom paling kiri
            tops = sorted({round(c['top'], 1) for c in chars
                           if c['x0'] < 30 and c['size'] > 8.8 and c['text'].isdigit() and c['top'] > 78})
            anchors = []
            for t in tops:
                if not anchors or t - anchors[-1] > 3:
                    anchors.append(t)

            small_l = _words(pg, lambda o: o['size'] < 8.6 and o['x1'] <= X_TOTAL_SPLIT)
            small_r = _words(pg, lambda o: o['size'] < 8.6 and o['x1'] > X_TOTAL_SPLIT)
            small = small_l + small_r
            by_x = sorted(chars, key=lambda c: c['x0'])

            for T in anchors:
                near = [c for c in by_x if abs(c['top'] - T) < 2 and c['size'] > 8.8]
                no = ''.join(c['text'] for c in near if c['x0'] < 30)
                nik = ''.join(c['text'] for c in near if 30 <= c['x0'] < X_NAME_MIN)
                name = ''
                started = False
                for c in near:
                    if not (X_NAME_MIN <= c['x0'] < X_NAME_MAX):
                        continue
                    if not c['text'].strip():
                        if started and c['x0'] >= X_NAME_JUNK_SPACE:
                            name += ' '
                        continue
                    started = True
                    name += c['text']
                name = re.sub(r'\s+', ' ', name).strip()

                line_a = [w for w in small if -2.5 <= w['top'] - T <= 4 and w['x0'] >= X_DAY_MIN]
                line_b = [w for w in small if 8 <= w['top'] - T <= 16 and w['x0'] >= X_DAY_MIN]

                codes = {}
                for w in line_a:
                    if w['x0'] < X_SUMMARY_MIN and not re.fullmatch(r'[\d.]+', w['text']):
                        d = min(day_x, key=lambda k: abs(day_x[k] - (w['x0'] + w['x1']) / 2))
                        codes[d] = w['text']
                svals = [float(w['text']) for w in sorted(
                    (w for w in line_a if w['x0'] >= X_SUMMARY_MIN), key=lambda w: w['x0'])]

                ot, ot_total = {}, None
                for w in line_b:
                    v = float(w['text'])
                    if w['x1'] > X_TOTAL_SPLIT:
                        ot_total = v
                    else:
                        d = min(day_x, key=lambda k: abs(day_x[k] - (w['x0'] + w['x1']) / 2))
                        ot[d] = v
                out.append(dict(Period=period, No=int(no), NIK=nik, Name=name,
                                svals=svals, ot_total=ot_total, codes=codes, ot=ot))
    return period, out


def transform(records):
    """Raw records -> (daily df, monthly df, qa-issue list)."""
    daily, summ, issues = [], [], []
    for r in records:
        y, m = int(r['Period'][:4]), int(r['Period'][4:])
        nd = calendar.monthrange(y, m)[1]
        sv = dict(zip(SUM_COLS, r['svals']))

        # validasi 1: total lembur
        if round(sum(r['ot'].get(d, 0) for d in range(1, 32)) - (r['ot_total'] or 0), 2) != 0:
            issues.append((r['Period'], r['NIK'], 'Jumlah OT harian != total di PDF'))
        # validasi 2: hitung kode vs ringkasan PDF
        cc = Counter(RAW_ALIAS.get(v, v) for v in r['codes'].values())
        for code in ['H', 'A', 'I', 'CT', 'S', 'SD', 'TL', 'IR', 'DL', 'CH', 'PC', 'N']:
            if cc.get(code, 0) != sv.get(code, 0):
                issues.append((r['Period'], r['NIK'], f'Hitungan {code}: harian={cc.get(code, 0)} vs PDF={sv.get(code)}'))

        for d in range(1, nd + 1):
            date = dt.date(y, m, d)
            raw = r['codes'].get(d)
            code = RAW_ALIAS.get(raw, raw)
            if code is None:
                code = 'OFF' if date.weekday() >= 5 else 'NODATA'
            daily.append([int(r['Period']), date, date.strftime('%a'), r['NIK'], r['Name'], raw, code,
                          STATUS.get(code, ('', 'Unknown'))[1],
                          0 if code in ('OFF', 'H', 'NODATA') else 1,
                          r['ot'].get(d, 0.0)])
        summ.append([int(r['Period']), r['No'], r['NIK'], r['Name']] +
                    [sv.get(c) for c in SUM_COLS] + [r['ot_total']])

    D = pd.DataFrame(daily, columns=['Period', 'Date', 'Day_Name', 'NIK', 'Name', 'Raw_Code', 'Status',
                                     'Status_Group', 'Is_Workday', 'OT_Hours'])
    S = pd.DataFrame(summ, columns=['Period', 'No', 'NIK', 'Name'] + SUM_COLS + ['OT_Total_Reported'])
    return D, S, issues


# ----------------------------------------------------------------------------
# Load -> Excel
# ----------------------------------------------------------------------------
def build_workbook(D, S):
    F = 'Arial'
    hfont = Font(name=F, bold=True, color='FFFFFF', size=10)
    hfill = PatternFill('solid', fgColor='1F3864')
    bfont = Font(name=F, size=10)

    def header(ws, n):
        for c in range(1, n + 1):
            cell = ws.cell(row=1, column=c)
            cell.font, cell.fill = hfont, hfill
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        ws.freeze_panes = 'A2'

    def body(ws, n):
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=n):
            for c in row:
                c.font = bfont

    def widths(ws, w):
        for i, x in enumerate(w, 1):
            ws.column_dimensions[get_column_letter(i)].width = x

    wb = Workbook()
    ws = wb.active
    ws.title = 'README'
    periods = sorted(D.Period.unique())
    notes = [
        ('HRIS Monthly Attendance - hasil konsolidasi ETL', True),
        (f'Periode: {periods[0]} s/d {periods[-1]} | dibuat {dt.datetime.now():%Y-%m-%d %H:%M}', False),
        ('', False),
        ('Sheet', True),
        ('Attendance_Daily : 1 baris per karyawan per tanggal (tabel fakta untuk dashboard)', False),
        ('Employee_Monthly : ringkasan per karyawan per bulan (angka PDF) + kolom cek berformula', False),
        ('Dim_Employee / Dim_Status : tabel dimensi', False),
        ('QA_Checks : rekap validasi', False),
        ('', False),
        ('Asumsi', True),
        ('1. Hanya hari kerja yang berkode di PDF; Sabtu/Minggu tanpa kode diisi Status = OFF.', False),
        ('2. Kode I* dinormalisasi ke IR (sesuai hitungan ringkasan HRIS); kode asli di Raw_Code.', False),
        ('3. Baris angka kedua per karyawan diasumsikan jam LEMBUR (OT_Hours).', False),
        ('4. Hari kerja tanpa kode diberi Status NODATA (misal karyawan keluar/masuk di tengah bulan).', False),
        ('5. Kolom CHK, HD, LC, CL4, CL5, CL6, SOR tidak punya legenda di PDF; disimpan apa adanya.', False),
    ]
    for i, (t, b) in enumerate(notes, 1):
        ws.cell(row=i, column=1, value=t).font = Font(name=F, bold=b, size=12 if i == 1 else 10)
    ws.column_dimensions['A'].width = 120

    # Attendance_Daily
    ws = wb.create_sheet('Attendance_Daily')
    ws.append(list(D.columns))
    for r in D.itertuples(index=False):
        ws.append([None if (isinstance(v, float) and pd.isna(v)) else v for v in r])
    nd = ws.max_row
    header(ws, D.shape[1]); body(ws, D.shape[1])
    for r in range(2, nd + 1):
        ws.cell(row=r, column=2).number_format = 'yyyy-mm-dd'
        ws.cell(row=r, column=10).number_format = '0.0'
    widths(ws, [10, 12, 10, 10, 30, 10, 9, 13, 11, 10])
    ws.auto_filter.ref = f'A1:J{nd}'

    # Employee_Monthly (+ formula check)
    ws = wb.create_sheet('Employee_Monthly')
    hdr = list(S.columns) + ['N_Calc', 'OT_Calc', 'Check_N', 'Check_OT']
    ws.append(hdr)
    ncs = S.shape[1]
    cN = get_column_letter(list(S.columns).index('N') + 1)
    cOT = get_column_letter(ncs)
    cNc, cOc = get_column_letter(ncs + 1), get_column_letter(ncs + 2)
    for i, r in enumerate(S.itertuples(index=False), 2):
        ws.append([None if (isinstance(v, float) and pd.isna(v)) else v for v in r] + [
            f'=COUNTIFS(Attendance_Daily!$A:$A,$A{i},Attendance_Daily!$D:$D,$C{i},Attendance_Daily!$G:$G,"N")',
            f'=SUMIFS(Attendance_Daily!$J:$J,Attendance_Daily!$A:$A,$A{i},Attendance_Daily!$D:$D,$C{i})',
            f'=IF({cNc}{i}={cN}{i},"OK","DIFF")',
            f'=IF(ROUND({cOc}{i}-{cOT}{i},2)=0,"OK","DIFF")'])
    ns = ws.max_row
    header(ws, len(hdr)); body(ws, len(hdr))
    widths(ws, [10, 6, 10, 30] + [6] * len(SUM_COLS) + [12, 9, 9, 9, 9])
    ws.auto_filter.ref = f'A1:{get_column_letter(len(hdr))}{ns}'
    ws.row_dimensions[1].height = 30

    # Dim_Employee
    ws = wb.create_sheet('Dim_Employee')
    ws.append(['NIK', 'Name', 'Periods_Count'])
    for e in D.groupby(['NIK', 'Name']).Period.nunique().reset_index().itertuples(index=False):
        ws.append(list(e))
    header(ws, 3); body(ws, 3); widths(ws, [10, 30, 14])
    ne = ws.max_row

    # Dim_Status
    ws = wb.create_sheet('Dim_Status')
    ws.append(['Status', 'Description', 'Status_Group'])
    for k, (d, g) in STATUS.items():
        ws.append([k, d, g])
    header(ws, 3); body(ws, 3); widths(ws, [10, 34, 14])

    # QA_Checks
    ws = wb.create_sheet('QA_Checks')
    ws.append(['Check', 'Result', 'Expected'])
    cCN, cCO = get_column_letter(ncs + 3), get_column_letter(ncs + 4)
    for q in [
        ['Baris Attendance_Daily', f'=COUNTA(Attendance_Daily!$A$2:$A${nd})', len(D)],
        ['Baris Employee_Monthly', f'=COUNTA(Employee_Monthly!$A$2:$A${ns})', len(S)],
        ['Karyawan unik', f'=COUNTA(Dim_Employee!$A$2:$A${ne})', D.NIK.nunique()],
        ['Check_N = DIFF', f'=COUNTIF(Employee_Monthly!${cCN}$2:${cCN}${ns},"DIFF")', 0],
        ['Check_OT = DIFF', f'=COUNTIF(Employee_Monthly!${cCO}$2:${cCO}${ns},"DIFF")', 0],
        ['Hari kerja NODATA', f'=COUNTIF(Attendance_Daily!$G$2:$G${nd},"NODATA")', int((D.Status == 'NODATA').sum())],
    ]:
        ws.append(q)
    header(ws, 3); body(ws, 3); widths(ws, [30, 14, 14])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ----------------------------------------------------------------------------
# Streamlit UI
# ----------------------------------------------------------------------------
@st.cache_data(show_spinner=False)
def run_etl(files):
    """files: tuple[(filename, bytes)] -> (D, S, issues, per_file_info)"""
    recs, info = [], []
    for fname, content in files:
        period, rows = parse_pdf(io.BytesIO(content))
        info.append(dict(File=fname, Period=period, Employees=len(rows)))
        recs += rows
    D, S, issues = transform(recs)
    return D, S, issues, info


def main():
    st.set_page_config(page_title='HRIS Attendance ETL', layout='wide')
    st.title('HRIS Attendance ETL')
    st.caption('PDF Monthly Attendance Report (PT.GLOSTAR INDONESIA) -> Excel siap dashboard')

    up = st.file_uploader('Upload PDF HRIS (boleh banyak bulan sekaligus)', type='pdf', accept_multiple_files=True)
    if not up:
        st.info('Upload satu atau beberapa PDF untuk memulai.')
        return

    files = tuple((f.name, f.getvalue()) for f in up)
    try:
        with st.spinner('Memproses PDF...'):
            D, S, issues, info = run_etl(files)
    except Exception as e:
        st.error(f'Gagal memproses PDF: {e}')
        return
    if D.empty:
        st.warning('Tidak ada data karyawan yang terbaca. Pastikan format PDF sama dengan laporan HRIS.')
        return

    dup = pd.DataFrame(info).Period.duplicated().any()
    if dup:
        st.warning('Ada lebih dari satu file dengan periode yang sama. Data akan terduplikasi.')

    # KPI
    w = D[D.Is_Workday == 1]
    c1, c2, c3, c4, c5 = st.columns(5)
    c1.metric('Periode', D.Period.nunique())
    c2.metric('Karyawan unik', D.NIK.nunique())
    c3.metric('Baris harian', f'{len(D):,}')
    c4.metric('Tingkat hadir', f'{(w.Status_Group == "Hadir").mean():.1%}' if len(w) else '-')
    c5.metric('Total jam lembur', f'{D.OT_Hours.sum():,.1f}')

    # QA
    if issues:
        st.error(f'{len(issues)} selisih validasi terhadap angka ringkasan PDF')
        st.dataframe(pd.DataFrame(issues, columns=['Period', 'NIK', 'Issue']), use_container_width=True)
    else:
        st.success('Validasi lulus: hitungan kode status dan total lembur cocok dengan ringkasan PDF.')
    nodata = D[D.Status == 'NODATA']
    if len(nodata):
        st.warning(f'{len(nodata)} hari kerja tanpa kode (NODATA): '
                   + ', '.join(sorted(nodata.Name.unique())))

    t1, t2, t3, t4 = st.tabs(['Attendance_Daily', 'Employee_Monthly', 'Ringkasan status', 'File terbaca'])
    with t1:
        per = st.multiselect('Filter periode', sorted(D.Period.unique()), default=sorted(D.Period.unique()))
        st.dataframe(D[D.Period.isin(per)], use_container_width=True, height=400)
    with t2:
        st.dataframe(S, use_container_width=True, height=400)
    with t3:
        piv = w.pivot_table(index='Period', columns='Status', values='NIK', aggfunc='count', fill_value=0)
        st.dataframe(piv, use_container_width=True)
        st.bar_chart(w.pivot_table(index='Period', columns='Status_Group', values='NIK',
                                   aggfunc='count', fill_value=0))
    with t4:
        st.dataframe(pd.DataFrame(info), use_container_width=True)

    st.divider()
    xlsx = build_workbook(D, S)
    stamp = f'{D.Period.min()}-{D.Period.max()}'
    d1, d2 = st.columns(2)
    d1.download_button('Download Excel (semua sheet)', xlsx, f'HRIS_Attendance_{stamp}.xlsx',
                       'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', type='primary')
    d2.download_button('Download CSV (Attendance_Daily, untuk Looker Studio)',
                       D.to_csv(index=False).encode('utf-8'), f'Attendance_Daily_{stamp}.csv', 'text/csv')


if __name__ == '__main__':
    main()
